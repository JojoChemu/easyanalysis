import asyncio
from io import BytesIO
from pathlib import Path
from tempfile import TemporaryDirectory
import unittest

import pandas as pd
from fastapi import HTTPException, UploadFile
from openpyxl import load_workbook

from backend.services import processing as reports
from backend.services.arrears_difference import generate_loan_arrears_difference
from backend.services.report_input import read_report
from backend.routers.processing import loan_arrears_difference_endpoint


def fixture():
    return pd.DataFrame([{
        "Account Number": "001234", "Client Reference Id": "001",
        "Client Name": "Customer", "Branch": "Nairobi Branch",
        "Loan Officer": "JANE DOE", "LoanOfficerName": "JANE DOE",
        "Disbursement Date": "2026-10-16", "Creation Date": "2026-10-01",
        "Expiry Date": "2026-11-16", "Total Loan Amount": "10,000",
        "Loan Amount": "10,000", "Outstanding Principal": "5,000",
        "Total In Arrears": "1,000", "PortfolioAtRisk(PAR)": "1,000",
        "Phone Number": "", "National Id Number": "",
        "Number Of Dependents": "", "IsPersonWithDisability": "No",
        "ImpairmentType": "", "MembershipPolicy": "", "Group": "",
        "Repayment Date": "2026-10-16", "Transaction Number": "001",
        "PrincipalAmountRepaid": "100", "InterestAmountRepaid": "10",
        "Charges": "0", "TotalAmountRepaid": "110",
    }])


def workbook_values(path):
    wb = load_workbook(path)
    return {ws.title: list(ws.values) for ws in wb}


class ReportTests(unittest.TestCase):
    def setUp(self):
        self.temp = TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.plain = self.root / "plain.csv"
        self.titled = self.root / "titled.csv"
        fixture().to_csv(self.plain, index=False)
        self.titled.write_text("All Loan Report" + "," * (len(fixture().columns) - 1) + "\n" + self.plain.read_text())
        self.targets = self.root / "Targets.xlsx"
        pd.DataFrame({"Loan Officer": ["JANE DOE", "PETER SMITH", "NEW OFFICER", "NO_ARREARS"], "Target": [20000] * 4, "RO CODE": [1, 2, 3, 4]}).to_excel(self.targets, index=False)

    def test_title_detection_csv_excel_and_preserved_identifiers(self):
        for path in (self.plain, self.titled):
            self.assertEqual(read_report(path, dtype={"Account Number": str}).iloc[0]["Account Number"], "001234")
        excel = self.root / "repayments.xlsx"
        with pd.ExcelWriter(excel) as writer:
            fixture().to_excel(writer, index=False, startrow=1)
            writer.sheets["Sheet1"]["A1"] = "Repayment Report"
        pd.testing.assert_frame_equal(read_report(excel), pd.read_excel(excel, skiprows=1))
        # Existing labels must never be mistaken for a title.
        pd.testing.assert_frame_equal(read_report(self.plain), pd.read_csv(self.plain))
        ambiguous = self.root / "ambiguous.csv"
        ambiguous.write_text("Unknown title\nnot,labels\n")
        self.assertIn("Unknown title", read_report(ambiguous).columns)

    def test_all_existing_workflows_with_and_without_titles(self):
        workflows = [
            (reports.clean_all_loans, False, ".csv"),
            (reports.clean_portfolio_risk, False, ".csv"),
            (reports.clean_loan_arrears, False, ".csv"),
            (reports.clean_repayment_report, False, ".xlsx"),
            (reports.generate_ro_dashboard, True, ".xlsx"),
            (reports.generate_ro_monthly_customer_mix, True, ".xlsx"),
            (reports.generate_par_performance, True, ".xlsx"),
        ]
        target_bytes = self.targets.read_bytes()
        for function, targets, suffix in workflows:
            with self.subTest(workflow=function.__name__):
                outputs = []
                for index, source in enumerate((self.plain, self.titled)):
                    if function == reports.generate_par_performance:
                        source = self.root / f"par-{index}.csv"
                        fixture().drop(columns="Loan Officer").to_csv(source, index=False)
                        if index:
                            source.write_text("Portfolio Risk Report\n" + source.read_text())
                    output = self.root / f"{function.__name__}-{index}{suffix}"
                    function(source, self.targets, output) if targets else function(source, output)
                    outputs.append(output)
                if suffix == ".csv":
                    self.assertEqual(outputs[0].read_bytes(), outputs[1].read_bytes())
                else:
                    self.assertEqual(workbook_values(outputs[0]), workbook_values(outputs[1]))
                    if function == reports.clean_repayment_report:
                        self.assertEqual(outputs[0].with_suffix('.csv').read_bytes(), outputs[1].with_suffix('.csv').read_bytes())
        self.assertEqual(target_bytes, self.targets.read_bytes())

    def arrears_sources(self):
        first = self.root / "Arrears-2026-10-01.csv"
        second = self.root / "Arrears-2026-10-16.csv"
        # Name order, totals, unmatched names, currency, and zero baselines.
        first.write_text('Loan Arrears Report,\nLoanOfficerName,Total In Arrears\nDOE JANE,"KES 1,000"\nPETER SMITH,200\nTOTAL,1200\nZZZZZZ,500\n')
        second.write_text('LoanOfficerName,Total In Arrears\nJANE DOE,500\nPETER SMITH,300\nNEW OFFICER,100\nGRAND TOTAL,900\n')
        return first, second

    def test_difference_calculations_ordering_and_workbook(self):
        first, second = self.arrears_sources()
        for index, sources in enumerate(((first, second), (second, first))):
            output = self.root / f"difference{index}.xlsx"
            generate_loan_arrears_difference(*sources, self.targets, output)
            wb = load_workbook(output)
            self.assertEqual(wb.sheetnames, ["Arrears Comparison", "Beginning Arrears", "MTD Arrears", "Executive Summary"])
            ws = wb["Arrears Comparison"]
            rows = {row[1]: row for row in list(ws.values)[5:]}
            self.assertEqual(rows["JANE DOE"], (1, "JANE DOE", 1000, 500, -500, -.5, "REDUCED"))
            self.assertEqual(rows["PETER SMITH"][2:], (200, 300, 100, .5, "INCREASED"))
            self.assertEqual(rows["NEW OFFICER"][2:], (0, 100, 100, 1, "INCREASED"))
            self.assertEqual(rows["NO_ARREARS"][2:], (0, 0, 0, 0, "NO CHANGE"))
            self.assertEqual(ws.freeze_panes, "A6")
            self.assertEqual(ws["F6"].number_format, "0.00%")
            self.assertEqual(ws["B2"].value, "01-Oct-2026")
            self.assertEqual(ws["E2"].value, "16-Oct-2026")

    def test_upload_endpoint_and_invalid_dates(self):
        first, second = self.arrears_sources()
        def upload(path, name=None):
            return UploadFile(filename=name or path.name, file=BytesIO(path.read_bytes()))
        response = asyncio.run(loan_arrears_difference_endpoint(upload(second), upload(first), upload(self.targets)))
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.body[:2], b"PK")
        self.assertIn("Arrears_Report.xlsx", response.headers["content-disposition"])
        for name in ["missing-date.csv", "Arrears-2026-10-01.csv", "Arrears-2026-99-01.csv"]:
            with self.subTest(filename=name), self.assertRaises(HTTPException) as error:
                asyncio.run(loan_arrears_difference_endpoint(upload(first), upload(second, name), upload(self.targets)))
            self.assertEqual(error.exception.status_code, 422)


if __name__ == "__main__":
    unittest.main()
