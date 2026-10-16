"""LoanArrearsDiff notebook adapted to uploaded file paths.

Calculations, matching threshold, date ordering, ranking, workbook sheets and
formatting follow the supplied notebook. Colab upload/display calls are omitted.
"""
from datetime import datetime, timedelta, timezone
from pathlib import Path
import re
import numpy as np
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils import get_column_letter
from rapidfuzz import process, fuzz
from backend.services.report_input import read_report
from backend.services.processing import ProcessingError

def generate_loan_arrears_difference(first_path: Path, second_path: Path, target_path: Path, output_path: Path):
    try:
        _generate(first_path, second_path, target_path, output_path)
    except (ValueError, KeyError, TypeError, IndexError) as error:
        raise ProcessingError(str(error)) from error

def _generate(first_path, second_path, target_path, output_path):
    csv_files = [str(first_path), str(second_path)]
    arrears_1 = read_report(first_path)
    arrears_2 = read_report(second_path)
    targets = pd.read_excel(target_path)

    def clean_column_names(df):
        df = df.copy()
        df.columns = df.columns.astype(str).str.strip().str.upper().str.replace('\\s+', ' ', regex=True)
        return df
    arrears_1 = clean_column_names(arrears_1)
    arrears_2 = clean_column_names(arrears_2)
    targets = clean_column_names(targets)

    def find_column(df, possible_names):
        normalized = {re.sub('[^A-Z0-9]', '', str(col).upper()): col for col in df.columns}
        for name in possible_names:
            key = re.sub('[^A-Z0-9]', '', name.upper())
            if key in normalized:
                return normalized[key]
        for col in df.columns:
            normalized_col = re.sub('[^A-Z0-9]', '', str(col).upper())
            for name in possible_names:
                normalized_name = re.sub('[^A-Z0-9]', '', name.upper())
                if normalized_name in normalized_col or normalized_col in normalized_name:
                    return col
        return None
    loan_officer_candidates = ['LOAN OFFICER', 'LOAN OFFICER NAME', 'LOANOFFICER', 'LOANOFFICERNAME', 'OFFICER', 'OFFICER NAME']
    arrears_candidates = ['TOTAL IN ARREARS', 'TOTAL ARREARS', 'ARREARS', 'AMOUNT IN ARREARS', 'ARREARS AMOUNT']
    target_candidates = ['TARGET', 'MONTHLY TARGET', 'MTD TARGET', 'ARREARS TARGET', 'COLLECTION TARGET']
    loan_officer_col_1 = find_column(arrears_1, loan_officer_candidates)
    loan_officer_col_2 = find_column(arrears_2, loan_officer_candidates)
    arrears_col_1 = find_column(arrears_1, arrears_candidates)
    arrears_col_2 = find_column(arrears_2, arrears_candidates)
    target_officer_col = find_column(targets, loan_officer_candidates)
    target_col = find_column(targets, target_candidates)
    required_columns = {'Arrears File 1 - Loan Officer': loan_officer_col_1, 'Arrears File 1 - Arrears': arrears_col_1, 'Arrears File 2 - Loan Officer': loan_officer_col_2, 'Arrears File 2 - Arrears': arrears_col_2, 'Targets - Loan Officer': target_officer_col, 'Targets - Target': target_col}
    missing = [name for name, column in required_columns.items() if column is None]
    if missing:
        raise ValueError('Please check the column names in your uploaded files.')

    def clean_amount(series):
        return series.astype(str).str.replace(',', '', regex=False).str.replace('KSh', '', regex=False).str.replace('KES', '', regex=False).str.replace(' ', '', regex=False).str.replace('(', '-', regex=False).str.replace(')', '', regex=False).replace(['', 'nan', 'None', 'NaN'], np.nan).astype(float).fillna(0)
    arrears_1[arrears_col_1] = clean_amount(arrears_1[arrears_col_1])
    arrears_2[arrears_col_2] = clean_amount(arrears_2[arrears_col_2])
    targets[target_col] = clean_amount(targets[target_col])

    def clean_names(series):
        return series.astype(str).str.strip().str.replace('\\s+', ' ', regex=True).str.upper()
    arrears_1[loan_officer_col_1] = clean_names(arrears_1[loan_officer_col_1])
    arrears_2[loan_officer_col_2] = clean_names(arrears_2[loan_officer_col_2])
    targets[target_officer_col] = clean_names(targets[target_officer_col])

    def remove_total_rows(df, officer_column):
        df = df.copy()
        total_names = ['TOTAL', 'TOTALS', 'GRAND TOTAL', 'GRANDTOTAL']
        mask = df[officer_column].astype(str).str.upper().str.strip().isin(total_names)
        return df.loc[~mask].copy()
    arrears_1 = remove_total_rows(arrears_1, loan_officer_col_1)
    arrears_2 = remove_total_rows(arrears_2, loan_officer_col_2)
    targets = remove_total_rows(targets, target_officer_col)

    def extract_date_from_filename(filename):
        name = Path(filename).name.replace('_', '-')
        match = re.search('(20\\d{2})-(\\d{1,2})-(\\d{1,2})', name)
        if match:
            year = int(match.group(1))
            month = int(match.group(2))
            day = int(match.group(3))
            try:
                return datetime(year, month, day)
            except ValueError:
                return None
        return None
    file_dates = {}
    for file in csv_files:
        detected_date = extract_date_from_filename(file)
        file_dates[file] = detected_date
    if any((date is None for date in file_dates.values())):
        raise ValueError('Both CSV filenames must contain a date such as 2026-8-1 or 2026-08-19.')
    if len(set(file_dates.values())) != 2:
        raise ValueError('Choose two arrears reports with different dates in their filenames.')
    sorted_files = sorted(file_dates.items(), key=lambda x: x[1])
    beginning_file = sorted_files[0][0]
    beginning_date = sorted_files[0][1]
    current_file = sorted_files[1][0]
    current_date = sorted_files[1][1]
    if beginning_file == csv_files[0]:
        beginning_arrears = arrears_1
        beginning_officer_col = loan_officer_col_1
        beginning_arrears_col = arrears_col_1
    else:
        beginning_arrears = arrears_2
        beginning_officer_col = loan_officer_col_2
        beginning_arrears_col = arrears_col_2
    if current_file == csv_files[0]:
        current_arrears = arrears_1
        current_officer_col = loan_officer_col_1
        current_arrears_col = arrears_col_1
    else:
        current_arrears = arrears_2
        current_officer_col = loan_officer_col_2
        current_arrears_col = arrears_col_2
    target_summary = targets[[target_officer_col, target_col]].copy()
    target_summary.columns = ['Loan Officer', 'Target']
    target_names = target_summary['Loan Officer'].dropna().astype(str).str.strip().unique().tolist()

    def match_name_to_target(arrears_name, target_names, threshold=60):
        if pd.isna(arrears_name):
            return None
        arrears_name = str(arrears_name).strip().upper()
        if arrears_name in target_names:
            return arrears_name
        match = process.extractOne(arrears_name, target_names, scorer=fuzz.token_sort_ratio)
        if match is not None:
            matched_name = match[0]
            score = match[1]
            if score >= threshold:
                return matched_name
        return None
    beginning_arrears = beginning_arrears.copy()
    beginning_arrears['Matched Loan Officer'] = beginning_arrears[beginning_officer_col].apply(lambda x: match_name_to_target(x, target_names))
    current_arrears = current_arrears.copy()
    current_arrears['Matched Loan Officer'] = current_arrears[current_officer_col].apply(lambda x: match_name_to_target(x, target_names))
    beginning_summary = beginning_arrears[beginning_arrears['Matched Loan Officer'].notna()].groupby('Matched Loan Officer', as_index=False)[beginning_arrears_col].sum()
    beginning_summary.columns = ['Loan Officer', 'Beginning Arrears']
    current_summary = current_arrears[current_arrears['Matched Loan Officer'].notna()].groupby('Matched Loan Officer', as_index=False)[current_arrears_col].sum()
    current_summary.columns = ['Loan Officer', 'MTD Arrears']
    report = target_summary[['Loan Officer']].copy()
    report = pd.merge(report, beginning_summary, on='Loan Officer', how='left')
    report = pd.merge(report, current_summary, on='Loan Officer', how='left')
    report['Beginning Arrears'] = report['Beginning Arrears'].fillna(0)
    report['MTD Arrears'] = report['MTD Arrears'].fillna(0)
    report['Change in Arrears'] = report['MTD Arrears'] - report['Beginning Arrears']
    report['% Change'] = np.where(report['Beginning Arrears'] != 0, report['Change in Arrears'] / report['Beginning Arrears'] * 100, np.where(report['MTD Arrears'] > 0, 100, 0))
    report['Status'] = np.where(report['Change in Arrears'] < 0, 'REDUCED', np.where(report['Change in Arrears'] > 0, 'INCREASED', 'NO CHANGE'))
    report['Rank'] = report['Change in Arrears'].rank(method='min', ascending=True).astype(int)
    report = report.sort_values(by='Change in Arrears', ascending=True).reset_index(drop=True)
    report = report[['Rank', 'Loan Officer', 'Beginning Arrears', 'MTD Arrears', 'Change in Arrears', '% Change', 'Status']]
    output_file = output_path
    with pd.ExcelWriter(output_file, engine='openpyxl') as writer:
        report.to_excel(writer, sheet_name='Arrears Comparison', index=False)
        beginning_summary.to_excel(writer, sheet_name='Beginning Arrears', index=False)
        current_summary.to_excel(writer, sheet_name='MTD Arrears', index=False)
    wb = load_workbook(output_file)
    dark_blue = '1F4E78'
    white = 'FFFFFF'
    green = 'C6EFCE'
    dark_green = '006100'
    red = 'FFC7CE'
    dark_red = '9C0006'
    ws = wb['Arrears Comparison']
    ws.insert_rows(1, 4)
    ws['A1'] = 'ARREARS PERFORMANCE REPORT'
    ws['A1'].font = Font(bold=True, size=18, color=white)
    ws['A1'].fill = PatternFill('solid', fgColor=dark_blue)
    ws['A1'].alignment = Alignment(horizontal='center', vertical='center')
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=7)
    ws.row_dimensions[1].height = 30
    ws['A2'] = 'Beginning-of-Month Date'
    ws['B2'] = beginning_date.strftime('%d-%b-%Y')
    ws['D2'] = 'MTD Report Date'
    ws['E2'] = current_date.strftime('%d-%b-%Y')
    ws['A3'] = 'Report Generated'
    ws['B3'] = datetime.now(timezone(timedelta(hours=3))).strftime('%d-%b-%Y %H:%M')
    for cell in ['A2', 'D2', 'A3']:
        ws[cell].font = Font(bold=True, color=white)
        ws[cell].fill = PatternFill('solid', fgColor=dark_blue)
    header_row = 5
    for cell in ws[header_row]:
        cell.font = Font(bold=True, color=white)
        cell.fill = PatternFill('solid', fgColor=dark_blue)
        cell.alignment = Alignment(horizontal='center', vertical='center')
    last_row = ws.max_row
    last_col = ws.max_column
    table_ref = f'A{header_row}:{get_column_letter(last_col)}{last_row}'
    tab = Table(displayName='ArrearsComparisonTable', ref=table_ref)
    style = TableStyleInfo(name='TableStyleMedium2', showFirstColumn=False, showLastColumn=False, showRowStripes=True, showColumnStripes=False)
    tab.tableStyleInfo = style
    ws.add_table(tab)
    for row in range(header_row + 1, last_row + 1):
        ws.cell(row, 3).number_format = '#,##0.00'
        ws.cell(row, 4).number_format = '#,##0.00'
        ws.cell(row, 5).number_format = '#,##0.00'
        ws.cell(row, 6).number_format = '0.00%'
        if isinstance(ws.cell(row, 6).value, (int, float)):
            ws.cell(row, 6).value /= 100
    for row in range(header_row + 1, last_row + 1):
        status = ws.cell(row, 7)
        if status.value == 'REDUCED':
            status.fill = PatternFill('solid', fgColor=green)
            status.font = Font(bold=True, color=dark_green)
        elif status.value == 'INCREASED':
            status.fill = PatternFill('solid', fgColor=red)
            status.font = Font(bold=True, color=dark_red)
    ws.conditional_formatting.add(f'E{header_row + 1}:E{last_row}', ColorScaleRule(start_type='min', start_color='63BE7B', mid_type='percentile', mid_value=50, mid_color='FFEB84', end_type='max', end_color='F8696B'))
    ws.freeze_panes = 'A6'
    widths = {'A': 10, 'B': 30, 'C': 22, 'D': 22, 'E': 22, 'F': 15, 'G': 18}
    for column, width in widths.items():
        ws.column_dimensions[column].width = width
    for sheet_name in ['Beginning Arrears', 'MTD Arrears']:
        sheet = wb[sheet_name]
        sheet.freeze_panes = 'A2'
        for cell in sheet[1]:
            cell.font = Font(bold=True, color=white)
            cell.fill = PatternFill('solid', fgColor=dark_blue)
            cell.alignment = Alignment(horizontal='center')
        for column_cells in sheet.columns:
            max_length = 0
            column_letter = get_column_letter(column_cells[0].column)
            for cell in column_cells:
                try:
                    max_length = max(max_length, len(str(cell.value)))
                except (TypeError, ValueError):
                    pass
            sheet.column_dimensions[column_letter].width = min(max_length + 3, 35)
    wb.save(output_file)
    total_beginning = report['Beginning Arrears'].sum()
    total_mtd = report['MTD Arrears'].sum()
    total_change = total_mtd - total_beginning
    if total_beginning != 0:
        total_percentage_change = total_change / total_beginning
    else:
        total_percentage_change = 0
    officers_reduced = (report['Status'] == 'REDUCED').sum()
    officers_increased = (report['Status'] == 'INCREASED').sum()
    officers_no_change = (report['Status'] == 'NO CHANGE').sum()
    summary = pd.DataFrame({'Measure': ['Beginning Arrears', 'MTD Arrears', 'Overall Change', 'Overall % Change', 'Officers Reduced', 'Officers Increased', 'Officers No Change'], 'Value': [total_beginning, total_mtd, total_change, total_percentage_change, officers_reduced, officers_increased, officers_no_change]})
    with pd.ExcelWriter(output_path, engine='openpyxl', mode='a') as writer:
        summary.to_excel(writer, sheet_name='Executive Summary', index=False)
